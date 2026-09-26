"""Parses the resolved bank (.docx): one answer letter and justification per question.

Expected layout, repeated 2737 times:
    N. <stem>
    Respuesta: x) <text of the correct option>
    Por qué: <justification>
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import docx

QUESTION_RE = re.compile(r"^(\d+)\.\s*(.*)$", re.S)
ANSWER_RE = re.compile(r"^Respuesta:\s*([a-h])\)\s*(.*)$", re.S)
WHY_RE = re.compile(r"^Por qué:\s*(.*)$", re.S)


@dataclass
class ResolvedAnswer:
    number: int
    stem: str
    letter: str
    answer_text: str
    justification: str


def parse_paragraphs(paragraphs: list[str]) -> list[ResolvedAnswer]:
    """Parses plain paragraph texts. Question numbers must be consecutive."""
    results: list[ResolvedAnswer] = []
    current: dict | None = None
    section = None

    def close():
        if current and current.get("letter"):
            results.append(
                ResolvedAnswer(
                    number=current["number"],
                    stem=current["stem"].strip(),
                    letter=current["letter"],
                    answer_text=current["answer"].strip(),
                    justification=current.get("why", "").strip(),
                )
            )

    for raw in paragraphs:
        text = raw.strip()
        if not text:
            continue
        expected = (current["number"] + 1) if current else 1
        q = QUESTION_RE.match(text)
        if q and int(q.group(1)) == expected and (current is None or section == "why"):
            close()
            current = {"number": expected, "stem": q.group(2)}
            section = "stem"
            continue
        if current is None:
            continue
        a = ANSWER_RE.match(text)
        if a and section == "stem":
            current["letter"], current["answer"] = a.group(1), a.group(2)
            section = "answer"
            continue
        w = WHY_RE.match(text)
        if w and section == "answer":
            current["why"] = w.group(1)
            section = "why"
            continue
        key = {"stem": "stem", "answer": "answer", "why": "why"}[section]
        current[key] = f"{current.get(key, '')}\n{text}"
    close()
    return results


def parse(docx_path: Path) -> list[ResolvedAnswer]:
    document = docx.Document(str(docx_path))
    return parse_paragraphs([p.text for p in document.paragraphs])


_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def answer_images(docx_path: Path) -> dict[int, list[bytes]]:
    """Returns, per question number, the images placed between 'Respuesta:' and
    'Por qué:' (the picture of the correct option when options are images)."""
    document = docx.Document(str(docx_path))
    result: dict[int, list[bytes]] = {}
    number = 0
    in_answer = False
    for p in document.paragraphs:
        text = p.text.strip()
        q = QUESTION_RE.match(text)
        if q and int(q.group(1)) == number + 1 and not in_answer:
            number += 1
            continue
        if ANSWER_RE.match(text):
            in_answer = True
        elif WHY_RE.match(text):
            in_answer = False
            continue
        if in_answer:
            for blip in p._p.iter(_BLIP):
                part = document.part.related_parts[blip.get(_EMBED)]
                result.setdefault(number, []).append(part.blob)
    return result
